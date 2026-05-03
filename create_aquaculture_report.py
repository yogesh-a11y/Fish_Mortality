from pathlib import Path
import re

import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.shared import Inches, Pt, RGBColor


OUT_DOCX = Path("Aquaculture_Decision_Support_Report.docx")
ARCH_IMG = Path("system_architecture.png")


def draw_architecture(path: Path) -> None:
    fig, ax = plt.subplots(figsize=(11.5, 6.2))
    ax.set_xlim(0, 12)
    ax.set_ylim(0, 7)
    ax.axis("off")

    boxes = [
        (0.45, 4.55, 2.0, 1.15, "IoT Sensor Inputs\nTemperature, DO,\npH, Turbidity"),
        (3.0, 4.55, 2.0, 1.15, "Data Cleaning\nDatetime sorting\nLeakage removal"),
        (5.55, 4.55, 2.0, 1.15, "Feature Engineering\nLag, delta,\nrolling, time"),
        (8.1, 4.55, 2.0, 1.15, "Random Forest\nClassifier\nbalanced weights"),
        (5.55, 2.0, 2.0, 1.15, "Model Artifact\nmodel.pkl\nfeature list"),
        (8.1, 2.0, 2.0, 1.15, "Streamlit App\nlive input form\nrisk probability"),
        (10.45, 2.0, 1.4, 1.15, "Decision\nStable or\nAt Risk"),
    ]

    colors = ["#D9EDF7", "#E8F5E9", "#FFF3CD", "#FCE4EC", "#EDE7F6", "#E3F2FD", "#FFE0B2"]
    for (x, y, w, h, text), color in zip(boxes, colors):
        patch = FancyBboxPatch(
            (x, y),
            w,
            h,
            boxstyle="round,pad=0.03,rounding_size=0.08",
            linewidth=1.4,
            edgecolor="#34495E",
            facecolor=color,
        )
        ax.add_patch(patch)
        ax.text(x + w / 2, y + h / 2, text, ha="center", va="center", fontsize=10)

    arrows = [
        ((2.45, 5.12), (3.0, 5.12)),
        ((5.0, 5.12), (5.55, 5.12)),
        ((7.55, 5.12), (8.1, 5.12)),
        ((9.1, 4.55), (6.55, 3.15)),
        ((7.55, 2.58), (8.1, 2.58)),
        ((10.1, 2.58), (10.45, 2.58)),
    ]
    for start, end in arrows:
        ax.add_patch(
            FancyArrowPatch(
                start,
                end,
                arrowstyle="-|>",
                mutation_scale=16,
                linewidth=1.5,
                color="#34495E",
            )
        )

    ax.text(
        6,
        6.45,
        "Architecture of the t+1 Aquaculture Risk Decision Support System",
        ha="center",
        va="center",
        fontsize=15,
        weight="bold",
    )
    ax.text(
        6,
        0.9,
        "The system predicts future risk y(t+1), not the current label y(t).",
        ha="center",
        va="center",
        fontsize=11,
        color="#555555",
    )
    plt.tight_layout()
    fig.savefig(path, dpi=220, bbox_inches="tight")
    plt.close(fig)


def add_paragraphs(doc: Document, text: str) -> None:
    for raw in [p.strip() for p in text.strip().split("\n\n") if p.strip()]:
        paragraph = doc.add_paragraph(raw)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def word_count(text: str) -> int:
    return len(re.findall(r"\b\w+\b", text))


report_sections = [
    (
        "Abstract",
        """
This project presents an AI-powered decision support system for aquaculture that predicts whether fish are likely to be at risk in the next one hour using water quality parameters. The central purpose of the work is not to classify the present condition of the fish pond, but to provide an early warning that supports preventive action before water quality deterioration becomes severe. The system uses temperature, dissolved oxygen, pH, turbidity, and derived temporal features to learn patterns associated with the next-hour health status. The target variable is therefore defined as target_t1, created by shifting the original Health Status column by one future time step. In practical terms, the model learns the relationship between the sensor condition at time t and the fish-risk state at time t+1.

The project follows a clean machine learning pipeline suitable for demonstration and further deployment. Data preprocessing includes column-name cleaning, datetime conversion, chronological sorting, and removal of leakage features such as survival rate, corrective measures, oxygenation interventions, thermal risk indices, and alert columns. These features are excluded because they either describe outcomes after an event or encode rule-based decisions that would not be available as independent real-time inputs. Feature engineering is a critical part of the system. Lag features represent the previous state of the pond, delta features measure short-term change, rolling statistics summarize recent dissolved oxygen behavior, and cyclical time encoding captures hourly patterns without treating time as a simple linear number.

The main model is a Random Forest Classifier trained with class_weight='balanced' to address the unequal distribution of Stable and At Risk cases. Logistic Regression is used as a transparent baseline. Evaluation uses a stratified train-test split, confusion matrix, precision, recall, F1-score, ROC-AUC, and PR-AUC. On the realistic simulated aquaculture dataset, the Random Forest achieved approximately 89.61 percent accuracy, with an At Risk recall of 45.77 percent and an At Risk F1-score of 58.82 percent. These results show that accuracy alone is not sufficient for aquaculture risk assessment because missing a true risk case can be more harmful than issuing a manageable false warning. The final system is implemented as a Streamlit application that loads the saved model artifact, accepts live water-quality readings, computes the required features, and outputs a Stable or At Risk prediction with risk probability.
""",
    ),
    (
        "Introduction",
        """
Aquaculture has become an important part of global food production, rural livelihood, and protein supply. As fish farming grows, farmers are expected to maintain productivity while also protecting animal welfare, water quality, and environmental sustainability. According to the Food and Agriculture Organization, modern aquaculture is central to the wider blue transformation agenda, where aquatic foods are treated as an important component of food security and sustainable production (Food and Agriculture Organization of the United Nations [FAO], 2024). However, aquaculture is highly sensitive to water quality. Fish live directly inside the production environment, so changes in dissolved oxygen, temperature, pH, turbidity, and related conditions can quickly affect feeding, growth, disease susceptibility, and mortality.

Traditional aquaculture management often depends on manual observation, periodic water testing, and farmer experience. These methods remain useful, but they can be too slow for rapidly changing pond conditions. Dissolved oxygen, for example, may decline during the night or early morning, especially when biological oxygen demand is high. Temperature influences metabolism and oxygen solubility. pH can shift due to photosynthesis, respiration, and water chemistry. Turbidity can increase after rainfall, feeding, or sediment disturbance. Because these variables interact, risk may not be visible from a single parameter alone. Boyd and Tucker (1998) emphasized that water quality is one of the major environmental concerns in aquaculture because the culture system itself can influence nutrient load, organic matter, and oxygen demand.

The rise of Internet of Things monitoring, low-cost sensors, and machine learning provides a practical opportunity to improve aquaculture decision support. IoT devices can continuously capture water-quality readings, while machine learning models can learn patterns from historical data and produce warnings. Recent studies have explored machine learning and IoT for aquaculture water quality monitoring and prediction, showing that data-driven systems can support faster and more consistent decisions than manual monitoring alone (Baena-Navarro et al., 2025; Li et al., 2022; Zambrano et al., 2021). This project follows the same direction but focuses specifically on future risk prediction. The goal is not simply to say whether the pond is currently stable. Instead, the goal is to predict whether the fish will be at risk in the next hour.

This distinction is important. A static classifier that predicts the current condition may only confirm a problem after it is already present. A one-hour-ahead early warning system gives the farmer time to inspect aeration, adjust water exchange, reduce feeding, or prepare corrective action. Even a short forecast horizon can be valuable in aquaculture because water quality interventions often need to happen quickly. Therefore, the project is framed as a real-time decision support system, where sensor readings are transformed into actionable risk information.
""",
    ),
    (
        "Problem Statement",
        """
Fish mortality and stress in aquaculture are often linked to rapid deterioration in water quality. Farmers may not always detect these changes early enough, especially when monitoring is manual, infrequent, or dependent on visual observation. A pond may appear normal while dissolved oxygen is declining, temperature is approaching stressful levels, or multiple moderate stressors are combining. By the time visible fish distress is observed, the response window may already be narrow.

The problem addressed in this project is the lack of a lightweight, interpretable, and demonstration-ready system that predicts next-hour fish risk from water quality parameters. The system must avoid a common machine learning mistake in applied projects: predicting the present label using features that would not be available in real time or using post-event information that leaks the answer into the model. Features such as survival rate, corrective measures, oxygenation interventions, low oxygen alerts, and thermal risk indices may be useful for reporting, but they should not be used as independent predictors for real-time risk forecasting. Including them would produce inflated performance and would not represent a scientifically valid early warning system.

The project therefore asks the following practical question: given current and recent water-quality readings, can a machine learning model predict whether fish will be At Risk in the next hour? The answer is developed through a pipeline that cleans the data, removes leakage features, engineers temporal predictors, shifts the target forward by one hour, trains baseline and Random Forest models, evaluates the model using risk-sensitive metrics, and deploys the trained model in a Streamlit application.
""",
    ),
    (
        "Objectives",
        """
The main objective of this project is to design and implement an AI-powered decision support system that predicts next-hour fish risk in aquaculture using water quality parameters. The specific objectives are as follows.

First, the project aims to prepare a clean dataset by standardizing column names, converting the Datetime column into a proper datetime format, sorting observations chronologically, and removing columns that could cause data leakage. Second, it aims to create temporal features that reflect realistic pond dynamics, including lag features, rolling statistics, delta features, and cyclical hour encoding. Third, it aims to define the target variable correctly as a future label, target_t1, so that the model predicts y(t+1) from features at time t rather than predicting the current state. Fourth, it aims to train and compare a Logistic Regression baseline with a Random Forest Classifier, using class weights to handle class imbalance without using synthetic oversampling. Fifth, it aims to evaluate model performance using metrics that are meaningful for aquaculture risk, especially recall and F1-score for the At Risk class. Sixth, it aims to save the trained model and feature list as a reusable artifact and implement a Streamlit application for live demonstration.

Together, these objectives support a complete workflow from data preparation to model deployment. The project is intentionally designed to be compact and understandable while still following machine learning practices expected in a master's-level applied research project.
""",
    ),
    (
        "Literature Review",
        """
Aquaculture water quality management has long been recognized as a core factor in fish health and production success. Boyd and Tucker (1998) described how aquaculture production systems can alter their own water quality through feeding, fertilization, waste accumulation, and biological oxygen demand. Their work remains important because it explains why parameters such as dissolved oxygen, temperature, pH, and turbidity are not isolated measurements. They are connected parts of a living production environment. Dissolved oxygen is especially important because fish require oxygen for respiration, and low oxygen conditions can cause stress, poor feeding, and mortality. Temperature affects oxygen solubility and fish metabolism, while pH and turbidity influence physiological comfort, visibility, and ecological balance.

Recent literature increasingly connects water quality monitoring with IoT and machine learning. Zambrano et al. (2021) studied machine learning for manually measured water quality prediction in fish farming and highlighted the importance of monitoring variables such as dissolved oxygen, pH, and pond temperature. Li et al. (2022) examined machine learning approaches for predicting aquaculture water quality and reported that models such as support vector machines can be effective for water-quality estimation. More recent work has integrated continuous monitoring and IoT-based sensing with machine learning models, including Random Forest methods, to support real-time aquaculture management (Baena-Navarro et al., 2025). Reviews of IoT sensors for aquaculture also emphasize that real-time data collection can reduce monitoring delays and support more timely decisions (Velasquez et al., 2025).

From a machine learning perspective, this project builds on established classification methods. Logistic Regression is used as a baseline because it is simple, interpretable, and useful for checking whether linear relationships are sufficient. Random Forest is used as the main model because it can capture nonlinear interactions between water-quality parameters without requiring strong assumptions about feature distributions. Breiman (2001) introduced Random Forests as an ensemble of decision trees that improves predictive performance by combining many randomized trees. Later discussions of Random Forests have emphasized their usefulness for tabular data, robustness, and ability to estimate feature importance (Schonlau & Zou, 2020).

Evaluation methodology is also important in this project. Accuracy can be misleading when the At Risk class is less frequent than the Stable class. A model could achieve high accuracy by mostly predicting Stable while missing many risk events. Precision, recall, and F1-score therefore provide a more meaningful view of model behavior. Recall is especially important because a false negative means the system failed to warn the farmer about an upcoming risk condition. Precision-recall analysis is commonly recommended when the positive class is relatively rare or operationally important (Davis & Goadrich, 2006). ROC-AUC is also useful for summarizing threshold-independent discrimination, although PR-AUC can be more informative when the minority class is the primary concern (Fawcett, 2006).

Class imbalance is another recurring issue in risk prediction. SMOTE was introduced as a synthetic minority oversampling method (Chawla et al., 2002), but synthetic oversampling is not always appropriate for temporal sensor data. In time-dependent data, generated samples may blur chronological structure, create unrealistic transitions, or make training data appear more diverse than the real sensor process. For this reason, this project avoids SMOTE and instead uses class_weight='balanced', allowing the model to place more weight on minority At Risk samples while preserving the original temporal records.
""",
    ),
    (
        "Proposed System",
        """
The proposed system is a real-time aquaculture risk decision support system that receives water-quality readings, transforms them into a model-ready feature vector, and predicts whether fish will be At Risk in the next hour. The system is built around a supervised binary classification model with two output classes: Stable and At Risk. The important design decision is that the output refers to the future time step, not the present time step. Therefore, the system functions as an early warning tool rather than a descriptive dashboard.

The system begins with data input. In a complete deployment, sensor readings would come from IoT devices installed in the pond or tank. For the demonstration version, the Streamlit interface accepts manual inputs for temperature, dissolved oxygen, pH, and turbidity. The training dataset also includes related weather context variables, but the deployed model focuses on the core real-time water-quality parameters and engineered temporal features. The raw Datetime value is not used directly as a feature because timestamps as raw numbers can mislead the model and do not generalize well. Instead, hour is converted into sine and cosine components to represent the daily cycle smoothly.

After input, the system applies preprocessing and feature engineering. Lag features represent the previous reading, delta features represent the difference between current and previous readings, rolling features summarize short-term dissolved oxygen behavior, and cyclical features represent time-of-day patterns. These features are then arranged in exactly the same order as the feature list saved during training. The trained Random Forest model outputs a probability for the At Risk class. A threshold, usually 0.5 for the basic demonstration, converts that probability into the final label. The Streamlit app displays the prediction, probability percentage, and warning message when risk is elevated.

The proposed system is deliberately lightweight. It avoids unnecessary complexity and focuses on a pipeline that can be explained, reproduced, and demonstrated. At the same time, it follows production-oriented habits: the model is saved as a pickle artifact, feature columns are stored with the model, and the app loads the artifact rather than retraining each time. This makes the system suitable for a master's-level demonstration and provides a foundation for future IoT integration.
""",
    ),
    (
        "System Architecture",
        """
The architecture of the system follows a simple end-to-end flow. Sensor readings enter the system, preprocessing prepares the data, feature engineering creates temporal predictors, the trained Random Forest model estimates next-hour risk, and the Streamlit application presents the prediction to the user. The architecture separates training-time work from prediction-time work. During training, the dataset is cleaned, the target is shifted, models are trained, and the selected model is saved. During prediction, the saved model is loaded and used only for inference.

This separation is important because a real decision support system must behave consistently after deployment. The Streamlit app must not create different feature names or train a new model with different assumptions. For this reason, the feature-column list is saved together with the model. At inference time, the app builds a DataFrame with the required columns and sends it to the classifier. If previous sensor readings are available, lag and rolling features should be computed from real sensor history. In the demonstration app, when no previous readings are available, neutral placeholders are used and the limitation is clearly explained to the user.
""",
    ),
    (
        "Methodology",
        """
The methodology follows a structured applied machine learning pipeline: data description, preprocessing, feature engineering, target creation, data splitting, imbalance handling, model training, evaluation, model saving, and application deployment. The pipeline is designed around scientific validity. It avoids leakage features, avoids using the current target as an input, avoids raw Datetime as a model feature, and defines the prediction task as future risk rather than static classification.

The first methodological decision is to sort all observations by Datetime. Temporal order matters because lag, delta, and rolling features depend on previous observations. If data were not sorted, the lag feature might represent a random previous row rather than the previous hour. The second methodological decision is to remove post-event and rule-based columns. This makes the task harder but more realistic. The third methodological decision is to shift the target one row upward using Health_Status.shift(-1). As a result, the features at time t are paired with the label at time t+1. This creates the early warning formulation required by the project objective.

The model is trained using a stratified 80-20 train-test split. Although time-based validation is often preferred for final forecasting studies, stratification is used here to preserve the minority At Risk proportion in both train and test sets. This is important because the At Risk class is less frequent, and an unstratified split could accidentally produce a test set with too few risk cases for meaningful evaluation. The final evaluation still discusses temporal limitations and recommends time-based validation as future work.
""",
    ),
    (
        "Dataset Description",
        """
The dataset used for the main report is aquaculture_realistic_dataset.xlsx. It contains 4,383 hourly records and ten columns. The main columns are Datetime, Temperature (Celsius), Dissolved Oxygen (mg/L), pH, Turbidity (NTU), Precipitation (inches), Average Temperature (Celsius), High Temperature (Celsius), Low Temperature (Celsius), and Health Status. The Health Status label contains two classes: Stable and At Risk. In the dataset used for evaluation, there are 3,672 Stable observations and 711 At Risk observations before feature engineering. After creating lag and rolling features and shifting the target, 4,380 usable rows remain, with 3,669 Stable and 711 At Risk labels in the future target.

The dataset is best described as a realistic simulated aquaculture IoT dataset. It was generated to mimic hourly water-quality behavior using temporal patterns, sensor relationships, random noise, and multi-parameter risk rules. Temperature was simulated with daily and seasonal variation plus autocorrelated noise. Dissolved oxygen was simulated with time-of-day variation and an inverse relationship with temperature, reflecting the fact that warmer water generally holds less oxygen. pH was generated with a relationship to dissolved oxygen, while turbidity was influenced by precipitation and previous turbidity behavior. Health Status was created from combined stress conditions rather than from direct fish mortality observations.

The main features used by the final model are the real-time water-quality parameters and engineered temporal features. Temperature helps represent thermal stress and metabolic pressure. Dissolved oxygen is a critical indicator because oxygen depletion can rapidly affect fish survival. pH reflects chemical balance and can influence physiological stress. Turbidity reflects suspended particles and can indicate rainfall effects, sediment disturbance, or water-quality deterioration. These variables are common in aquaculture monitoring and have direct relevance to fish welfare.

The dataset has two important issues. The first is imbalance: Stable cases are much more common than At Risk cases. This is realistic because risk events are usually less frequent than normal operating periods. The second issue is potential leakage in broader aquaculture datasets. Some datasets may include survival rate, corrective interventions, oxygenation actions, thermal risk indices, or alert labels. These columns should not be used as predictive features because they either summarize the result or represent decisions made after the condition is already known. In this project, such columns are removed when present.
""",
    ),
    (
        "Data Preprocessing",
        """
Data preprocessing begins by loading the Excel dataset and cleaning column names. Spaces are replaced with underscores, units are simplified, and names are standardized so that the model pipeline can refer to columns consistently. For example, Temperature (Celsius) is converted to Temperature_C, Dissolved Oxygen (mg/L) is converted to Dissolved_Oxygen_mg_L, and Turbidity (NTU) is converted to Turbidity_NTU. This may appear minor, but consistent naming reduces errors in production code and makes the pipeline easier to maintain.

The Datetime column is converted using pandas datetime parsing and invalid timestamps are removed. The dataset is then sorted by Datetime. Sorting is essential because temporal feature engineering depends on correct chronological order. If the records are out of order, the model may receive lag and rolling features that do not represent true previous values. After sorting, the index is reset to keep row order clean.

Leakage feature removal is one of the most important preprocessing steps. Any feature that directly or indirectly reveals the current or future target must be excluded. Survival rate, low oxygen alerts, thermal risk index, corrective measures, oxygenation interventions, oxygenation automatic labels, and similar rule-based or post-event columns are excluded when present. The reason is simple: a real-time prediction system would not know these values before making its prediction. If they are used in training, the model may appear highly accurate, but it would be learning from information that is unavailable at prediction time. This would produce misleading results and reduce scientific validity.

Raw Datetime is also excluded from the model inputs. A timestamp as a raw value may encourage the model to memorize calendar position rather than learn meaningful water-quality patterns. However, time still matters in aquaculture, especially because oxygen, temperature, and pH can follow daily cycles. Therefore, hour is extracted from Datetime and encoded cyclically using sine and cosine transformations.
""",
    ),
    (
        "Feature Engineering",
        """
Feature engineering is critical in this project because the prediction problem is temporal. The model should not only see the current values of temperature, dissolved oxygen, pH, and turbidity; it should also receive information about recent changes and short-term history. In aquaculture, the direction and speed of change can be as important as the absolute value. A dissolved oxygen value may still be acceptable, but if it is falling quickly, the pond may become risky within the next hour.

Lag features are created for dissolved oxygen, temperature, and pH. These are DO_lag_1, Temp_lag_1, and pH_lag_1. A lag feature represents the value from the previous time step, which in this dataset corresponds to the previous hour. Domain-wise, this is useful because water quality is autocorrelated. The current state of a pond is usually related to its recent state. Fish stress risk can be influenced by sustained low oxygen or sustained high temperature rather than by a single isolated reading.

Delta features measure change from the previous hour. DO_delta is computed as current dissolved oxygen minus previous dissolved oxygen, and Temp_delta is computed as current temperature minus previous temperature. These features help the model understand trend direction. A negative DO_delta indicates that oxygen is decreasing. A positive Temp_delta indicates warming. In pond management, these changes can be early signals even before values cross a fixed threshold.

Rolling statistics are used to summarize short-term behavior. DO_mean_3 is the mean dissolved oxygen across the current and previous two readings, while DO_std_3 measures variability across the same three-hour window. Dissolved oxygen is selected for rolling features because it is one of the most immediate fish-risk parameters. A low rolling mean may show persistent oxygen stress, while higher rolling variability may indicate unstable water conditions. Rolling features help the model avoid overreacting to one noisy reading while still capturing short-term patterns.

Cyclical time encoding is created using hour_sin and hour_cos. This converts the hour of day into two smooth features using sin(2*pi*hour/24) and cos(2*pi*hour/24). The reason for using sine and cosine instead of raw hour is that time is circular. Hour 23 and hour 0 are close in real life, but as raw numbers they look far apart. Cyclical encoding preserves this structure and allows the model to learn daily patterns such as overnight oxygen decline or daytime photosynthesis effects.

The final feature set includes Temperature_C, Dissolved_Oxygen_mg_L, pH, Turbidity_NTU, DO_lag_1, Temp_lag_1, pH_lag_1, DO_delta, Temp_delta, DO_mean_3, DO_std_3, hour_sin, and hour_cos. The current Health_Status label is not included as a feature. This is essential because the model must learn from sensor data, not from the answer it is supposed to predict.
""",
    ),
    (
        "Target Definition",
        """
The target definition is the core of the project. The goal is not to classify the current condition of the fish pond. The goal is to predict whether fish will be At Risk in the next hour. To achieve this, the original Health Status column is converted into a binary variable, where Stable is 0 and At Risk is 1. Then a shifted target is created:

target_t1 = Health_Status shifted by -1.

This means that the label for each row is taken from the next time step. The features at time t are paired with the health status at time t+1. In notation, the model learns X(t) -> y(t+1). The final row is dropped because it has no future label after shifting. Rows with missing lag or rolling values are also dropped.

This target design converts the project from a static classifier into an early warning system. A static classifier might answer, "Is the fish pond currently risky?" The proposed system answers, "Given the current and recent water-quality conditions, is the fish pond likely to be risky in the next hour?" This distinction is scientifically and operationally important. Farmers need enough time to respond. A one-hour warning may allow them to inspect aerators, check oxygenation, adjust feeding, perform water exchange, or prepare corrective measures. The model therefore supports decision making rather than simply describing the present.
""",
    ),
    (
        "Train-Test Split",
        """
The dataset is split using train_test_split with test_size=0.2, random_state=42, and stratify=y. The result is an 80 percent training set and a 20 percent test set, while preserving the proportion of Stable and At Risk cases in both subsets. Stratification is justified because the target is imbalanced. If a random split were used without stratification, the test set might contain too few At Risk samples, making recall and F1-score unreliable.

In pure forecasting applications, a chronological time-based split is often preferred because it better tests future generalization. However, for this project demonstration, stratified splitting provides stable class representation and allows fair comparison between baseline and main models. The limitation is acknowledged in the report, and time-based validation is recommended as future work before real farm deployment.
""",
    ),
    (
        "Imbalance Handling",
        """
The At Risk class is less frequent than the Stable class. This reflects real monitoring conditions, where normal operating periods usually occur more often than risk events. However, class imbalance can cause a model to favor the majority class. A model that predicts Stable most of the time may achieve high accuracy while failing to detect the events that matter most.

This project handles imbalance using class_weight='balanced'. This option automatically gives higher weight to the minority class during training, encouraging the model to pay more attention to At Risk samples. The approach is simple, reproducible, and does not create artificial records.

SMOTE is intentionally avoided. Although SMOTE is a well-known method for creating synthetic minority samples (Chawla et al., 2002), it can be problematic for temporal water-quality data. Sensor readings are ordered and autocorrelated, meaning that the sequence matters. Synthetic samples inserted into the feature space may not represent physically realistic transitions between hours. They may also blur the relationship between current readings, lag values, deltas, and rolling statistics. For this reason, preserving the original temporal structure is preferred, and class weighting is used instead of oversampling.
""",
    ),
    (
        "Machine Learning Model: Random Forest Classifier",
        """
The main model used in the project is the Random Forest Classifier. A Random Forest is an ensemble learning method that trains many decision trees and combines their predictions. Each tree is trained using randomness in the data and feature selection, which reduces overfitting compared with a single decision tree. The final prediction is based on the aggregated output of all trees. In this project, the Random Forest is configured with n_estimators=200, class_weight='balanced', and random_state=42.

Random Forest is suitable for this project for several reasons. First, aquaculture sensor data is tabular, and Random Forest models often perform well on tabular datasets without requiring extensive scaling. Second, the relationship between water-quality parameters and fish risk is nonlinear. For example, a temperature value may become risky only when dissolved oxygen is also low, or pH may matter differently under different oxygen conditions. Random Forest can capture these interactions through tree splits. Third, it is relatively robust to noise and can handle mixed feature types after numeric encoding. Fourth, it provides feature importance estimates, which help interpret which variables influenced the model.

Logistic Regression is used as a baseline model. It is simpler and more interpretable, but it assumes a mostly linear relationship between features and the log-odds of risk. In the realistic dataset, Logistic Regression achieved lower accuracy than the tree-based models, suggesting that nonlinear relationships and feature interactions are important. An Extra Trees baseline was also evaluated as a random tree ensemble baseline, but Random Forest remains the selected main model because it is widely accepted, stable, and easy to explain in an applied decision-support setting.
""",
    ),
    (
        "Evaluation Metrics",
        """
The model is evaluated using confusion matrix, precision, recall, F1-score, ROC-AUC, and PR-AUC. The confusion matrix shows the number of true Stable predictions, false alarms, missed risk cases, and correctly detected risk cases. This is important because the cost of each error is not equal. A false positive may cause unnecessary attention or intervention, but a false negative may allow a dangerous condition to continue without warning.

Precision measures how many predicted At Risk cases are actually At Risk. Recall measures how many actual At Risk cases are detected by the model. F1-score balances precision and recall into a single value. For aquaculture, recall is especially critical. A high-recall system is less likely to miss fish-risk events, which is important when the goal is early warning and mortality prevention. However, recall must be balanced with precision so that farmers are not overwhelmed by excessive false alarms.

Accuracy is reported but interpreted carefully. Accuracy can be useful when classes are balanced, but in this project Stable cases are much more common than At Risk cases. Therefore, a high accuracy score can hide poor detection of the minority class. ROC-AUC is briefly considered as a threshold-independent measure of ranking ability, while PR-AUC is useful because the positive At Risk class is less frequent and operationally important (Davis & Goadrich, 2006; Fawcett, 2006).
""",
    ),
    (
        "Results and Discussion",
        """
On aquaculture_realistic_dataset.xlsx, the Random Forest model achieved an accuracy of 89.61 percent on the stratified test set. The confusion matrix was [[720, 14], [77, 65]], where 720 Stable cases were correctly classified, 14 Stable cases were falsely predicted as At Risk, 77 At Risk cases were missed, and 65 At Risk cases were correctly detected. The At Risk precision was 82.28 percent, At Risk recall was 45.77 percent, and At Risk F1-score was 58.82 percent. The ROC-AUC was approximately 0.784, and the PR-AUC was approximately 0.622.

These results require careful interpretation. The accuracy appears strong, but the At Risk recall shows that the model still misses a meaningful number of risk cases. In aquaculture decision support, this is important because the purpose of the system is not only to be correct on average. The purpose is to warn farmers before risky conditions develop. A model with high accuracy but low risk recall may look good statistically while still failing operationally. Therefore, future improvements should focus on improving recall, possibly by tuning the classification threshold, adding more sensor history, using time-based validation, or collecting real farm data with more diverse risk events.

The feature importance values provide useful insight. The most important features in the Random Forest were DO_mean_3, DO_lag_1, pH_lag_1, pH, Dissolved_Oxygen_mg_L, and Temperature_C. This makes domain sense. Dissolved oxygen and pH are important water-quality indicators, and their recent history helps represent whether the pond is moving toward stress. Temperature also contributes because it affects fish metabolism and oxygen solubility. The hour_sin and hour_cos features had lower importance, but they still help represent daily patterns.

The earlier dataset, Data_Model_IoTMLCQ_2024.xlsx, produced approximately 99.89 percent accuracy under the same clean t+1 pipeline. Such near-perfect performance should be discussed cautiously. In applied machine learning, extremely high accuracy can indicate that the dataset contains deterministic rules, duplicated patterns, static repeated values, or features that strongly encode the target. Even if obvious leakage columns are removed, the underlying target may still be generated from simple thresholds that the model can learn very easily. Therefore, the more realistic dataset result around 89-90 percent is more believable for academic discussion.

The results support the feasibility of a next-hour aquaculture risk prediction system, but they also show that accuracy is not the final goal. A practical system should be tuned for risk detection. If deployed in a real farm, the threshold could be lowered to increase recall, accepting more false alarms in exchange for fewer missed risk events. The appropriate threshold should be selected with farmers and aquaculture experts because the cost of false positives and false negatives depends on the production system.
""",
    ),
    (
        "System Implementation",
        """
The system is implemented as a lightweight Python machine learning pipeline and Streamlit application. The training pipeline loads the dataset, cleans column names, converts and sorts Datetime, removes leakage columns, engineers features, creates the shifted target, splits the data, trains models, evaluates performance, and saves the trained Random Forest model. The saved artifact is model.pkl. It contains both the model and the feature column list. Saving the feature list is important because the Streamlit app must send features to the model in the same structure used during training.

The Streamlit app provides a simple interface for demonstration. The user enters temperature, dissolved oxygen, pH, and turbidity. The app uses the current system time to compute hour_sin and hour_cos. In a full IoT deployment, lag and rolling features would come from stored previous sensor readings. In the standalone demonstration app, if previous readings are not available, the app uses neutral placeholders and clearly explains this limitation. For example, current dissolved oxygen may be used as DO_lag_1, and DO_delta may be set to zero. This allows the model to run while making the limitation transparent.

After constructing the input features, the app loads model.pkl using joblib and calls predict_proba to estimate the probability of the At Risk class. The app displays the predicted label, the risk probability as a percentage, and a warning message if the risk is high. This makes the output understandable to non-technical users. A farmer or demonstrator does not need to inspect raw model outputs; they receive a decision-support message that can trigger monitoring or intervention.

The implementation is production-oriented in a basic sense because training and inference are separated. The app does not retrain the model every time it runs. The feature list is preserved with the model. The code avoids hidden leakage features. The output is probability-based, allowing future threshold tuning. These choices make the project suitable for demonstration and provide a clear path toward integration with real sensor systems.
""",
    ),
    (
        "Software Requirements",
        """
The project is implemented in Python. The main libraries are pandas for data loading and preprocessing, numpy for numerical operations and cyclical encoding, scikit-learn for machine learning models and evaluation metrics, joblib for saving and loading the trained model, openpyxl for reading Excel files, matplotlib and seaborn for visualization, and Streamlit for the web application. These dependencies are listed in requirements.txt so that the project can be reproduced in another environment.

The recommended software environment includes Python 3.10 or later, Jupyter Notebook for research and explanation, and a local terminal for running the Streamlit app. The app can be launched with streamlit run streamlit_app.py. Hardware requirements are modest because the dataset is small and the Random Forest model is not computationally heavy. A standard laptop is sufficient for training and demonstration. For real deployment, additional components would be needed, such as sensor hardware, microcontroller or gateway devices, local storage or cloud database, and networking for real-time data transfer.
""",
    ),
    (
        "Limitations",
        """
The project has several limitations that should be stated clearly. The first limitation is that the main dataset is realistic simulated data rather than direct measurements from a working fish farm. It was designed to imitate plausible aquaculture sensor behavior, but it cannot capture all environmental, biological, and operational variation found in real ponds. Real farms may experience sensor calibration errors, missing data, fouling, sudden equipment failures, disease outbreaks, feeding events, rainfall shocks, and management actions that are not fully represented in the simulated dataset.

The second limitation is that the target label is rule-based rather than based on observed fish mortality events. Health Status was generated using stress conditions and label noise. This makes the dataset useful for model development and demonstration, but it does not prove that the model can predict actual mortality in all aquaculture settings. A real mortality prediction system would require labeled farm data linking sensor histories to observed stress or mortality outcomes.

The third limitation is temporal validation. The project uses stratified train-test splitting to preserve class balance, but a final forecasting model should also be tested using chronological splits. A chronological split would train on earlier dates and test on later dates, better representing real deployment. The fourth limitation is the handling of lag and rolling features in the Streamlit app. In a real system, these features require previous sensor readings. The standalone demo uses placeholders when history is unavailable, which is acceptable for demonstration but not ideal for production.

The fifth limitation is generalization. The model is trained on one dataset with specific ranges and simulated relationships. Different fish species, pond designs, climates, stocking densities, feeding regimes, and sensor types may produce different patterns. Before deployment, the model should be retrained or calibrated using local farm data.
""",
    ),
    (
        "Future Work",
        """
Future work should focus on improving realism, recall, interpretability, and deployment readiness. The first improvement is to collect real IoT sensor data from aquaculture ponds or tanks. Real data would allow the model to learn from actual environmental variability and observed fish outcomes. The second improvement is to implement multi-step forecasting, such as t+3 prediction, so that the system can warn about risk three hours ahead. A longer horizon may be valuable for planning, although it will likely be harder to predict accurately.

The third improvement is threshold tuning. Since recall is critical, the default probability threshold of 0.5 may not be optimal. The threshold could be lowered to detect more At Risk cases, with the trade-off of increased false alarms. The best threshold should be selected based on farm tolerance for false alarms and the cost of missed warnings. The fourth improvement is SHAP explainability. SHAP values can help explain why a specific prediction was made, such as whether low dissolved oxygen, high temperature, or unstable pH contributed most to the warning (Lundberg & Lee, 2017).

The fifth improvement is deployment integration. A real system should connect to sensors, store recent readings, compute lag and rolling features from actual history, and send alerts through a dashboard, SMS, or mobile notification. The model should also be monitored after deployment to detect performance drift. If seasonal patterns, sensor calibration, or farming practices change, the model may need retraining. Finally, future versions could compare Random Forest with gradient boosting, recurrent models, or temporal deep learning approaches, but only after a strong baseline and reliable data collection process are established.
""",
    ),
    (
        "Conclusion",
        """
This project developed a complete AI-powered decision support system for aquaculture risk prediction using water-quality parameters. The system is designed around future prediction, not static classification. By creating target_t1 from the next-hour Health Status label, the model learns to predict y(t+1) from X(t). This makes the system an early warning tool that can support preventive action in fish farming.

The project includes a clean data pipeline, leakage removal, temporal feature engineering, class imbalance handling, Random Forest model development, evaluation, model saving, and Streamlit deployment. Feature engineering plays a central role because aquaculture water quality is time-dependent. Lag features, delta features, rolling dissolved oxygen statistics, and cyclical time encoding allow the model to learn recent trends and daily patterns rather than relying only on current readings.

The Random Forest model achieved 89.61 percent accuracy on the realistic simulated dataset, with useful but improvable detection of At Risk cases. The result shows that the approach is feasible, but it also highlights why recall and F1-score are more important than accuracy alone. In aquaculture, a missed risk event can have serious biological and economic consequences. Therefore, future work should improve recall, test chronological validation, integrate real IoT data, and add explainability.

Overall, the project demonstrates a realistic and technically valid pathway for using machine learning in aquaculture decision support. It is compact enough for demonstration, transparent enough for academic reporting, and structured enough to be extended toward real farm deployment.
""",
    ),
]


references = [
    "Baena-Navarro, R., Carriazo-Regino, Y., Torres-Hoyos, F., & Pinedo-Lopez, J. (2025). Intelligent prediction and continuous monitoring of water quality in aquaculture: Integration of machine learning and Internet of Things for sustainable management. Water, 17(1), 82. https://doi.org/10.3390/w17010082",
    "Boyd, C. E., & Tucker, C. S. (1998). Pond aquaculture water quality management. Kluwer Academic Publishers. https://doi.org/10.1007/978-1-4615-5407-3",
    "Breiman, L. (2001). Random forests. Machine Learning, 45, 5-32. https://doi.org/10.1023/A:1010933404324",
    "Chawla, N. V., Bowyer, K. W., Hall, L. O., & Kegelmeyer, W. P. (2002). SMOTE: Synthetic minority over-sampling technique. Journal of Artificial Intelligence Research, 16, 321-357. https://doi.org/10.1613/jair.953",
    "Davis, J., & Goadrich, M. (2006). The relationship between precision-recall and ROC curves. Proceedings of the 23rd International Conference on Machine Learning, 233-240. https://doi.org/10.1145/1143844.1143874",
    "Essamlali, I., Nhaila, H., & El Khaili, M. (2024). Advances in machine learning and IoT for water quality monitoring: A comprehensive review. Heliyon, 10(6), e27920. https://doi.org/10.1016/j.heliyon.2024.e27920",
    "Fawcett, T. (2006). An introduction to ROC analysis. Pattern Recognition Letters, 27(8), 861-874. https://doi.org/10.1016/j.patrec.2005.10.010",
    "Food and Agriculture Organization of the United Nations. (2024). The state of world fisheries and aquaculture 2024: Blue transformation in action. FAO. https://www.fao.org/publications/fao-flagship-publications/the-state-of-world-fisheries-and-aquaculture",
    "Hastie, T., Tibshirani, R., & Friedman, J. (2009). The elements of statistical learning: Data mining, inference, and prediction (2nd ed.). Springer.",
    "Kuhn, M., & Johnson, K. (2013). Applied predictive modeling. Springer.",
    "Li, T., Lu, J., Wu, J., Zhang, Z., & Chen, L. (2022). Predicting aquaculture water quality using machine learning approaches. Water, 14(18), 2836. https://doi.org/10.3390/w14182836",
    "Lundberg, S. M., & Lee, S. I. (2017). A unified approach to interpreting model predictions. Advances in Neural Information Processing Systems, 30.",
    "Pedregosa, F., Varoquaux, G., Gramfort, A., Michel, V., Thirion, B., Grisel, O., Blondel, M., Prettenhofer, P., Weiss, R., Dubourg, V., Vanderplas, J., Passos, A., Cournapeau, D., Brucher, M., Perrot, M., & Duchesnay, E. (2011). Scikit-learn: Machine learning in Python. Journal of Machine Learning Research, 12, 2825-2830.",
    "Schonlau, M., & Zou, R. Y. (2020). The random forest algorithm for statistical learning. The Stata Journal, 20(1), 3-29. https://doi.org/10.1177/1536867X20909688",
    "Streamlit. (2026). Streamlit documentation. https://docs.streamlit.io/",
    "Velasquez, A., et al. (2025). Internet of Things (IoT) sensors for water quality monitoring in aquaculture systems: A systematic review and bibliometric analysis. AgriEngineering, 7(3), 78.",
    "Zambrano, A. F., Giraldo, L. F., Quimbayo, J., Medina, B., & Castillo, E. (2021). Machine learning for manually-measured water quality prediction in fish farming. PLoS ONE, 16(8), e0256380. https://doi.org/10.1371/journal.pone.0256380",
]


def build_docx() -> None:
    draw_architecture(ARCH_IMG)

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = doc.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(12)
    styles["Title"].font.name = "Times New Roman"
    styles["Title"].font.size = Pt(18)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("AI-Powered Decision Support System for Aquaculture")
    run.bold = True
    run.font.size = Pt(18)
    run.font.name = "Times New Roman"
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Real-Time Fish Mortality Risk Prediction Using Water Quality Parameters")
    run.bold = True
    run.font.size = Pt(14)
    run.font.name = "Times New Roman"
    doc.add_paragraph()

    overview = doc.add_paragraph()
    overview.alignment = WD_ALIGN_PARAGRAPH.CENTER
    overview.add_run("Prepared as a master's-level project report").italic = True
    doc.add_page_break()

    all_text = []
    for heading, content in report_sections:
        add_heading(doc, heading, 1)
        add_paragraphs(doc, content)
        all_text.append(content)
        if heading == "System Architecture":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            doc.add_picture(str(ARCH_IMG), width=Inches(6.6))
            caption = doc.add_paragraph("Figure 1. Proposed system architecture for next-hour aquaculture risk prediction.")
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption.runs[0].italic = True

        if heading == "Dataset Description":
            table = doc.add_table(rows=1, cols=3)
            table.style = "Table Grid"
            hdr = table.rows[0].cells
            hdr[0].text = "Feature"
            hdr[1].text = "Role"
            hdr[2].text = "Domain relevance"
            rows = [
                ("Temperature_C", "Current sensor input", "Thermal stress, metabolism, oxygen solubility"),
                ("Dissolved_Oxygen_mg_L", "Current sensor input", "Immediate oxygen availability for fish respiration"),
                ("pH", "Current sensor input", "Chemical balance and physiological comfort"),
                ("Turbidity_NTU", "Current sensor input", "Suspended particles, rainfall, sediment disturbance"),
                ("Lag and rolling features", "Temporal predictors", "Recent pond history and short-term trends"),
                ("hour_sin, hour_cos", "Cyclical time features", "Daily water-quality cycles without raw timestamp leakage"),
            ]
            for row in rows:
                cells = table.add_row().cells
                for i, value in enumerate(row):
                    cells[i].text = value

        if heading == "Results and Discussion":
            table = doc.add_table(rows=1, cols=2)
            table.style = "Table Grid"
            hdr = table.rows[0].cells
            hdr[0].text = "Metric"
            hdr[1].text = "Random Forest result"
            rows = [
                ("Accuracy", "89.61%"),
                ("At Risk precision", "82.28%"),
                ("At Risk recall", "45.77%"),
                ("At Risk F1-score", "58.82%"),
                ("ROC-AUC", "0.784"),
                ("PR-AUC", "0.622"),
                ("Confusion matrix", "[[720, 14], [77, 65]]"),
            ]
            for metric, value in rows:
                cells = table.add_row().cells
                cells[0].text = metric
                cells[1].text = value

    add_heading(doc, "References", 1)
    for ref in references:
        p = doc.add_paragraph(ref)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.left_indent = Inches(0.25)

    total_words = word_count("\n".join(all_text))
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = note.add_run(f"Approximate report body word count: {total_words}")
    run.italic = True
    run.font.color.rgb = RGBColor(90, 90, 90)

    doc.save(OUT_DOCX)
    print(f"Saved {OUT_DOCX.resolve()}")
    print(f"Saved {ARCH_IMG.resolve()}")
    print(f"Approximate body word count: {total_words}")


if __name__ == "__main__":
    build_docx()
